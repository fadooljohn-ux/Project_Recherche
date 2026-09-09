"""Public-dependency workbook closeout. Reads saved evidence; never runs science."""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

REPO = Path(__file__).resolve().parents[1]
MASTER = REPO / 'outputs/recherche-master-20260908/Project-Recherche-Master.xlsx'


def read(p):
    return json.loads(Path(p).read_text())


def digest(p):
    with Path(p).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def values(sheet):
    return list(sheet.values)


def append(sheet, row):
    n = sheet.max_row + 1
    for col, value in enumerate(row, 1):
        cell = sheet.cell(n, col, value)
        cell._style = copy(sheet.cell(n - 1, col)._style)
        cell.alignment = copy(sheet.cell(n - 1, col).alignment)
    sheet.row_dimensions[n].height = 90
    for table in sheet.tables.values():
        table.ref = f'A1:{get_column_letter(sheet.max_column)}{n}'
        if table.autoFilter:
            table.autoFilter.ref = table.ref


def add_work(sheet, row):
    old = next((r for r in values(sheet)[1:] if r[0] == row[0]), None)
    if old:
        if tuple(row[2:]) != old[2:len(row)]:
            raise ValueError('Existing work ID differs; use a new ID: ' + row[0])
        return
    append(sheet, row)


def apply(wb, root, mode, review):
    targets, searches, work = (wb[n] for n in ['Targets', 'Searches', 'Work log'])
    now = datetime.now().replace(microsecond=0)
    target_rows = {r[0]: i for i, r in enumerate(values(targets), 1)}
    if review:
        context = read(review)
        evidence = Path(context['evidence'])
        if not evidence.is_absolute():
            evidence = REPO / evidence
        sha = digest(evidence)
        for entry in context['entries']:
            add_work(work, [entry[0], now, *entry[1:], str(evidence), sha])
        for update in context.get('target_updates', []):
            row = target_rows[update['target']]
            for col, value in [(6, update['next_step']), (7, update['dataset_status']), (17, str(evidence)), (18, sha)]:
                targets.cell(row, col, value)
        return
    allowed = {'': 'MPTA', '--tpa-observed': 'TPA', '--ng15-observed': 'NG15',
               '--pta-observed': 'PTA', '--utmost-observed': 'UTMOST', '--ipta-observed': 'IPTA'}
    if mode not in allowed:
        raise ValueError('Fallback supports observed closeout and --review. Use a review context for preparation/diagnosis records.')
    selection_path = root / 'selection.json'
    selection = read(selection_path)
    freeze_path = root / 'execution-freeze.json'
    freeze = read(freeze_path)
    if freeze['selection_sha256'] != digest(selection_path):
        raise ValueError('Frozen selection changed')
    batch = root.parent.name + '-' + root.name if mode == '--pta-observed' else root.name
    for target in selection['selected']:
        row = target_rows[target]
        result_path = root / target / 'run01/result.json'
        if not result_path.exists():
            state = 'INCOMPLETE' if result_path.parent.exists() else 'NOT_LAUNCHED'
            add_work(work, [f'{batch}-{target}-{state.lower()}', now, 'Observed attempt', target, state,
                           'No terminal result exists; not a nondetection', 'Inspect retained evidence before continuing',
                           str(freeze_path), digest(freeze_path)])
            continue
        result = read(result_path)
        if result['target'] != target or result['execution_freeze_sha256'] != digest(freeze_path):
            raise ValueError('Result target or execution freeze differs')
        if digest(result_path.parent / 'periodogram.npz') != result['periodogram_sha256']:
            raise ValueError('Periodogram differs')
        record_id = f'{batch}-{target}-run01'
        old = next((r for r in values(searches)[1:] if r[0] == record_id or r[19] == str(result_path)), None)
        if old:
            if old[20] != digest(result_path):
                raise ValueError('Existing scientific evidence changed: ' + record_id)
            continue
        prep = read(root / target / 'prepared/receipt.json')
        bounds = prep.get('policy', selection.get('policy'))
        disposition = 'PERIODIC_SIGNAL_CANDIDATE' if result['candidate'] else result['status']
        detail = result['threshold_scope'] + '. ' + result['sensitivity_definition']
        if result['noise_adequacy_flag']:
            detail += ' Noise adequacy flagged.'
        if prep.get('checks', {}).get('historical_summary_difference_flag'):
            detail += ' Historical PAR summary differs; fixed-model likelihood verified, historical reproduction not claimed.'
        append(searches, [record_id, target, selection.get('dataset', allowed[mode]) + '; circular search', disposition,
                         result['peak_period_days'], 'Strongest eligible grid peak; candidate status is separate',
                         result['peak_amplitude_us'], result['peak_statistic'], result['threshold'], prep['toas'],
                         prep.get('epochs', prep.get('observing_days')), bounds['minimum_period_days'], bounds['maximum_period_days'],
                         result['grid_cells'], result['grid_cells'] - result['eligible_cells'],
                         result.get('sensitivity_median_worst_phase_us', result.get('sensitivity_median_worst_phase_amplitude_us')),
                         detail, None, 'Preserve consumed record; new data/scope requires a new ID', str(result_path), digest(result_path),
                         f"observed={result['observed_sha256']}; calibration={result['calibration_key']}"])
        prior_candidate = targets.cell(row, 3).value == 'PERIODIC_SIGNAL_CANDIDATE'
        if result['candidate'] or not prior_candidate:
            for col, value in [(3, disposition), (4, record_id), (5, 'Threshold crossing; review required' if result['candidate'] else 'No qualifying signal in the frozen eligible cells')]:
                targets.cell(row, col, value)
        targets.cell(row, 6, 'Preserve prior candidate review and assess this dataset separately' if prior_candidate else
                     'Bounded deep review and three-measure grading required' if result['candidate'] else 'Preserve result; new data or scope required for another search')
        targets.cell(row, 7, 'SEARCHED')
        add_work(work, [record_id + '-PORTABLE-CLOSEOUT', now, 'Observed closeout', target, disposition,
                       'Saved terminal result registered; no scientific execution by updater',
                       'Update source queue from reconciled campaign totals', str(result_path), digest(result_path)])
        if result['candidate']:
            append(wb['Candidate grades'], [record_id, target, 'REVIEW_PENDING', 'NOT_ASSESSED', 'NOT_ASSESSED',
                   None, None, None, None, None, 'Global FAP not established; bounded review required.',
                   None, None, 'REVIEW_PENDING', None, None])


def save_exact_numbers(wb, destination):
    """Keep Python float values through XLSX's default 16-digit serialization."""
    wb.save(destination)
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(destination) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    sheets = {f'xl/worksheets/sheet{i}.xml': sheet for i, sheet in enumerate(wb, 1)}
    with zipfile.ZipFile(destination, 'w') as archive:
        for item, content in members:
            if item.filename in sheets:
                root = ET.fromstring(content)
                for cell in root.findall('.//s:sheetData/s:row/s:c', ns):
                    value = sheets[item.filename][cell.attrib['r']].value
                    element = cell.find('s:v', ns)
                    if isinstance(value, float) and element is not None and cell.get('t', 'n') == 'n':
                        element.text = repr(value)
                content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            archive.writestr(item, content)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('root', nargs='?', type=Path)
    parser.add_argument('--review', type=Path)
    parser.add_argument('--master', type=Path, default=MASTER, help='Override only for an isolated development copy')
    args, flags = parser.parse_known_args()
    if len(flags) > 1 or not (args.root or args.review):
        parser.error('Supply a batch root and optional observed mode, or --review CONTEXT.json')
    master = args.master.resolve()
    lock = master.with_name(master.name + '.update-lock')
    descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    pending = None
    try:
        before_hash = digest(master)
        if args.review:
            evidence = Path(read(args.review)['evidence'])
            if not evidence.is_absolute(): evidence = REPO / evidence
            expected = read(evidence).get('master_before_sha256')
            if expected and expected != before_hash:
                raise ValueError('Master changed since review snapshot')
        wb = openpyxl.load_workbook(master)
        before = {s.title: values(s) for s in wb}
        apply(wb, args.root.resolve() if args.root else None, flags[0] if flags else '', args.review)
        for name in ['Searches', 'Work log', 'Candidate grades']:
            if values(wb[name])[:len(before[name])] != before[name]:
                raise ValueError('Historical records changed: ' + name)
        if [r[15] for r in values(wb['Targets'])[1:]] != [r[15] for r in before['Targets'][1:]]:
            raise ValueError('Operator notes changed')
        changed = any(values(wb[n]) != rows for n, rows in before.items())
        receipt = {'status': 'UNCHANGED', 'before_sha256': before_hash, 'after_sha256': before_hash,
                   'backend': 'openpyxl', 'scientific_runs_executed': 0,
                   'searches_added': wb['Searches'].max_row - len(before['Searches']),
                   'work_log_rows_added': wb['Work log'].max_row - len(before['Work log'])}
        if changed:
            history = REPO.parent / 'Recherche Recovery/20260907T223404/master-workbook-history' if master == MASTER.resolve() else master.parent / 'workbook-history'
            backup = history / (before_hash + '.xlsx')
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(master, backup)
            if digest(backup) != before_hash: raise ValueError('Backup differs')
            wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
            with tempfile.NamedTemporaryFile(dir=master.parent, suffix='.xlsx', delete=False) as f:
                pending = Path(f.name)
            save_exact_numbers(wb, pending)
            check = openpyxl.load_workbook(pending)
            for name in before:
                if values(check[name]) != values(wb[name]):
                    mismatch = next((i, j, a, b) for i, (left, right) in enumerate(zip(values(wb[name]), values(check[name])), 1) for j, (a, b) in enumerate(zip(left, right), 1) if a != b)
                    raise ValueError('Workbook round-trip differs: ' + name + ' ' + repr(mismatch))
            if digest(master) != before_hash: raise ValueError('Master changed during update')
            os.replace(pending, master)
            receipt.update(status='UPDATED', after_sha256=digest(master), backup=str(backup),
                           formulas='Preserved; Excel/LibreOffice recalculates on open; no cached-value claim')
        destination = args.review.resolve().parent if args.review else args.root.resolve()
        (destination / 'master-workbook-update.json').write_text(json.dumps(receipt, indent=2) + '\n')
        print(json.dumps(receipt, indent=2))
    finally:
        if pending and pending.exists(): pending.unlink()
        os.close(descriptor)
        lock.unlink()


if __name__ == '__main__':
    main()
